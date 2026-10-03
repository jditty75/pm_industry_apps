#!/usr/bin/env python3
"""
Download container-bound *_DM Deployment Health workbooks as XLSX via Drive export.

Read-only Google APIs only. Script/spreadsheet file IDs stay in gitignored artifacts only
(never written to tracked repo files). Uses ~/.clasprc.json for Script API + optional
gitignored Drive token when clasp lacks drive.readonly.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import webbrowser
from dataclasses import dataclass
from datetime import datetime, timezone
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path
from threading import Thread
from typing import Any

REPO_ROOT = Path(__file__).resolve().parents[3]
APPS_JSON = REPO_ROOT / "config" / "apps.json"
LIVE_DIR = REPO_ROOT / "docs" / "migrations" / "live"
ANALYSIS_DIR = REPO_ROOT / "docs" / "analysis" / "dm-family"
FRESHNESS_HOST = REPO_ROOT / "solutions" / "SLG_DM" / "src" / "DataFreshnessMonitorHost.js"
DRIVE_TOKEN_PATH = Path(__file__).resolve().parent / ".dm-snapshot-oauth.json"

DRIVE_READONLY_SCOPE = "https://www.googleapis.com/auth/drive.readonly"
EXPORT_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FRESHNESS_KEY_MAP = {
    "SLG_DM": "SLG",
    "HENP_DM": "HENP",
    "HC_DM": "HC",
    "EVI_DM": "EVI",
    "HS_DM": "AI",
}


@dataclass
class DmApp:
    app_id: str
    name: str
    project_path: str

    @property
    def clasp_path(self) -> Path:
        return REPO_ROOT / self.project_path.replace("/", "\\") / ".clasp.json"


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def git_sha() -> str:
    try:
        out = subprocess.check_output(
            ["git", "-C", str(REPO_ROOT), "rev-parse", "HEAD"],
            stderr=subprocess.DEVNULL,
            text=True,
        )
        return out.strip()
    except (subprocess.CalledProcessError, FileNotFoundError):
        return "unknown"


def discover_dm_apps() -> list[DmApp]:
    if not APPS_JSON.is_file():
        raise SystemExit(f"Missing {APPS_JSON}")
    data = load_json(APPS_JSON)
    apps: list[DmApp] = []
    for entry in data.get("applications", []):
        app_id = entry.get("appId", "")
        project_path = entry.get("projectPath", "")
        if not app_id.endswith("_DM"):
            continue
        if not project_path.startswith("solutions/"):
            continue
        full = REPO_ROOT / project_path
        if not full.is_dir():
            continue
        apps.append(
            DmApp(
                app_id=app_id,
                name=entry.get("name", app_id),
                project_path=project_path.replace("\\", "/"),
            )
        )
    apps.sort(key=lambda a: a.app_id)
    if not apps:
        raise SystemExit("No *_DM applications discovered.")
    return apps


def clasprc_path() -> Path:
    return Path.home() / ".clasprc.json"


def refresh_oauth_token(
    client_id: str,
    client_secret: str,
    refresh_token: str,
) -> dict[str, Any]:
    body = urllib.parse.urlencode(
        {
            "client_id": client_id,
            "client_secret": client_secret,
            "refresh_token": refresh_token,
            "grant_type": "refresh_token",
        }
    ).encode("utf-8")
    req = urllib.request.Request(
        "https://oauth2.googleapis.com/token",
        data=body,
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=60) as resp:
        return json.loads(resp.read().decode("utf-8"))


def get_clasprc_credentials() -> dict[str, str]:
    path = clasprc_path()
    if not path.is_file():
        raise SystemExit(
            "CLASP credentials not found. Run `clasp login` first, then retry."
        )
    tok = load_json(path)["tokens"]["default"]
    needed = ("client_id", "client_secret", "refresh_token")
    for key in needed:
        if not tok.get(key):
            raise SystemExit(f"Invalid {path}: missing {key}")
    refreshed = refresh_oauth_token(tok["client_id"], tok["client_secret"], tok["refresh_token"])
    return {
        "access_token": refreshed["access_token"],
        "client_id": tok["client_id"],
        "client_secret": tok["client_secret"],
        "refresh_token": tok["refresh_token"],
    }


def get_drive_credentials() -> dict[str, str]:
    if DRIVE_TOKEN_PATH.is_file():
        stored = load_json(DRIVE_TOKEN_PATH)
        tok = stored.get("token", stored)
        refreshed = refresh_oauth_token(
            tok["client_id"], tok["client_secret"], tok["refresh_token"]
        )
        return {
            "access_token": refreshed["access_token"],
            "source": "dm-snapshot-oauth",
        }
    creds = get_clasprc_credentials()
    return {"access_token": creds["access_token"], "source": "clasprc"}


def http_get_json(url: str, access_token: str) -> dict[str, Any]:
    req = urllib.request.Request(url, headers={"Authorization": f"Bearer {access_token}"})
    with urllib.request.urlopen(req, timeout=120) as resp:
        return json.loads(resp.read().decode("utf-8"))


def http_download(url: str, access_token: str) -> bytes:
    req = urllib.request.Request(url, headers={"Authorization": f"Bearer {access_token}"})
    with urllib.request.urlopen(req, timeout=600) as resp:
        return resp.read()


def read_script_id(app: DmApp) -> str:
    clasp = app.clasp_path
    if not clasp.is_file():
        raise FileNotFoundError(f"Missing local binding: {clasp}")
    data = load_json(clasp)
    script_id = data.get("scriptId")
    if not script_id:
        raise ValueError(f"No scriptId in {clasp}")
    return script_id


def resolve_parent_spreadsheet_id(script_id: str, access_token: str) -> str:
    url = f"https://script.googleapis.com/v1/projects/{script_id}"
    project = http_get_json(url, access_token)
    parent_id = project.get("parentId") or ""
    if not parent_id:
        raise ValueError("Apps Script project has no parentId (not container-bound?)")
    return parent_id


def drive_file_metadata(file_id: str, access_token: str) -> dict[str, Any]:
    fields = "modifiedTime,name,mimeType,version,size"
    q = urllib.parse.urlencode(
        {"fields": fields, "supportsAllDrives": "true"}
    )
    url = f"https://www.googleapis.com/drive/v3/files/{file_id}?{q}"
    return http_get_json(url, access_token)


def export_spreadsheet_xlsx(file_id: str, access_token: str) -> bytes:
    q = urllib.parse.urlencode(
        {
            "mimeType": EXPORT_MIME,
            "supportsAllDrives": "true",
        }
    )
    url = f"https://www.googleapis.com/drive/v3/files/{file_id}/export?{q}"
    try:
        return http_download(url, access_token)
    except urllib.error.HTTPError as err:
        body = err.read().decode("utf-8", errors="replace")
        if err.code in (403, 404):
            raise PermissionError(
                "Drive export denied (clasp token likely lacks drive.readonly). "
                "Run one of:\n"
                "  clasp login --extra-scopes https://www.googleapis.com/auth/drive.readonly\n"
                f"  python {Path(__file__).name} --authorize\n"
                f"Details: HTTP {err.code}"
            ) from err
        raise RuntimeError(f"Drive export failed HTTP {err.code}: {body[:500]}") from err


def parse_freshness_crosscheck() -> dict[str, str]:
    if not FRESHNESS_HOST.is_file():
        return {}
    text = FRESHNESS_HOST.read_text(encoding="utf-8-sig")
    block = re.search(
        r"configureDataFreshnessSpreadsheetIds\(\{([^}]+)\}\)",
        text,
        re.DOTALL,
    )
    if not block:
        return {}
    out: dict[str, str] = {}
    for key, val in re.findall(r"(\w+):\s*'([^']+)'", block.group(1)):
        out[key] = val
    return out


def freshness_warning(app_id: str, resolved_id: str, freshness: dict[str, str]) -> str | None:
    key = FRESHNESS_KEY_MAP.get(app_id)
    if not key or key not in freshness:
        return None
    legacy = freshness[key]
    if legacy != resolved_id:
        return (
            f"WARNING: DataFreshnessMonitorHost legacy map key {key} "
            "does not match Apps Script parentId (stale registry; ignored)."
        )
    return None


def write_gitignored_meta(path: Path, payload: dict[str, Any]) -> None:
    path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")


def snapshot_one(
    app: DmApp,
    *,
    what_if: bool,
    force: bool,
    script_token: str,
    drive_token: str,
    drive_auth_source: str,
    freshness: dict[str, str],
) -> dict[str, Any]:
    LIVE_DIR.mkdir(parents=True, exist_ok=True)
    xlsx_path = LIVE_DIR / f"{app.app_id}.xlsx"
    meta_path = LIVE_DIR / f"{app.app_id}.meta.json"

    result: dict[str, Any] = {
        "appId": app.app_id,
        "name": app.name,
        "projectPath": app.project_path,
        "xlsxFile": xlsx_path.name,
        "resolutionMethod": "apps_script_api_parentId",
        "driveAuthSource": drive_auth_source,
    }

    script_id = read_script_id(app)
    parent_id = resolve_parent_spreadsheet_id(script_id, script_token)
    warn = freshness_warning(app.app_id, parent_id, freshness)
    if warn:
        result["freshnessCrossCheck"] = warn
        print(warn, file=sys.stderr)

    try:
        meta = drive_file_metadata(parent_id, drive_token)
        result["driveModifiedTime"] = meta.get("modifiedTime")
        result["driveVersion"] = meta.get("version")
        result["driveMimeType"] = meta.get("mimeType")
    except urllib.error.HTTPError as err:
        result["driveMetadataError"] = f"HTTP {err.code}"
        meta = {}

    if what_if:
        result["action"] = "would_export"
        result["targetPath"] = str(xlsx_path.relative_to(REPO_ROOT)).replace("\\", "/")
        if xlsx_path.is_file():
            result["existingLocalFile"] = True
            result["existingLocalBytes"] = xlsx_path.stat().st_size
        return result

    if xlsx_path.is_file() and not force:
        age_h = (time.time() - xlsx_path.stat().st_mtime) / 3600
        result["action"] = "skipped_existing"
        result["reason"] = f"Use -Force to replace ({age_h:.1f}h old)"
        return result

    content = export_spreadsheet_xlsx(parent_id, drive_token)
    xlsx_path.write_bytes(content)

    meta_payload = {
        "appId": app.app_id,
        "exportedAtUtc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "driveModifiedTime": meta.get("modifiedTime"),
        "driveVersion": meta.get("version"),
        "resolutionMethod": "apps_script_api_parentId",
        "driveAuthSource": drive_auth_source,
        "gitShaAtExport": git_sha(),
        "xlsxFileName": xlsx_path.name,
        "xlsxBytes": len(content),
    }
    write_gitignored_meta(meta_path, meta_payload)

    result["action"] = "exported"
    result["xlsxBytes"] = len(content)
    return result


def run_authorize() -> int:
    creds = get_clasprc_credentials()
    port = 8765
    redirect_uri = f"http://127.0.0.1:{port}/"
    scope = urllib.parse.quote(DRIVE_READONLY_SCOPE, safe="")
    auth_url = (
        "https://accounts.google.com/o/oauth2/v2/auth?"
        + urllib.parse.urlencode(
            {
                "client_id": creds["client_id"],
                "redirect_uri": redirect_uri,
                "response_type": "code",
                "scope": DRIVE_READONLY_SCOPE,
                "access_type": "offline",
                "prompt": "consent",
            }
        )
    )

    code_holder: dict[str, str] = {}

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):  # noqa: N802
            parsed = urllib.parse.urlparse(self.path)
            params = urllib.parse.parse_qs(parsed.query)
            if "code" in params:
                code_holder["code"] = params["code"][0]
                self.send_response(200)
                self.end_headers()
                self.wfile.write(b"Authorization received. You can close this tab.")
            else:
                self.send_response(400)
                self.end_headers()
                self.wfile.write(b"Missing code.")

        def log_message(self, *_args):  # noqa: D102
            return

    server = HTTPServer(("127.0.0.1", port), Handler)
    thread = Thread(target=server.handle_request, daemon=True)
    thread.start()
    print("Opening browser for Drive read-only consent (Workday Google account)...")
    webbrowser.open(auth_url)
    thread.join(timeout=300)
    server.server_close()
    if "code" not in code_holder:
        print("Authorization timed out or failed.", file=sys.stderr)
        return 1

    body = urllib.parse.urlencode(
        {
            "client_id": creds["client_id"],
            "client_secret": creds["client_secret"],
            "code": code_holder["code"],
            "grant_type": "authorization_code",
            "redirect_uri": redirect_uri,
        }
    ).encode("utf-8")
    req = urllib.request.Request(
        "https://oauth2.googleapis.com/token", data=body, method="POST"
    )
    with urllib.request.urlopen(req, timeout=60) as resp:
        token_resp = json.loads(resp.read().decode("utf-8"))
    if not token_resp.get("refresh_token"):
        print(
            "No refresh_token returned. Revoke prior access or use prompt=consent again.",
            file=sys.stderr,
        )
        return 1

    DRIVE_TOKEN_PATH.write_text(
        json.dumps(
            {
                "token": {
                    "client_id": creds["client_id"],
                    "client_secret": creds["client_secret"],
                    "refresh_token": token_resp["refresh_token"],
                    "scope": DRIVE_READONLY_SCOPE,
                }
            },
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    print(f"Saved Drive read-only token to {DRIVE_TOKEN_PATH.name} (gitignored).")
    return 0


def analyze_workbook(app_id: str, xlsx_path: Path) -> dict[str, Any]:
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("openpyxl is required for structural analysis: pip install openpyxl")

    wb = openpyxl.load_workbook(xlsx_path, read_only=False, data_only=False)
    sheets_out = []
    formula_refs: list[dict[str, str]] = []
    ref_pattern = re.compile(
        r"(?:'([^']+)'|([A-Za-z0-9_]+))!(\$?[A-Z]{1,3}\$?\d+)"
    )

    for name in wb.sheetnames:
        ws = wb[name]
        state = ws.sheet_state
        hidden = state in ("hidden", "veryHidden")
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        headers: list[str] = []
        if max_row >= 1 and max_col >= 1:
            for col in range(1, min(max_col, 200) + 1):
                val = ws.cell(row=1, column=col).value
                if val is None:
                    continue
                if isinstance(val, str):
                    headers.append(val[:120])
                else:
                    headers.append(str(val)[:120])

        formula_count = 0
        sample_formulas = 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" or (
                    isinstance(cell.value, str) and cell.value.startswith("=")
                ):
                    formula_count += 1
                    if sample_formulas < 25:
                        ftxt = str(cell.value) if cell.value else ""
                        for m in ref_pattern.finditer(ftxt):
                            sheet_ref = m.group(1) or m.group(2)
                            formula_refs.append(
                                {
                                    "fromSheet": name,
                                    "fromCell": cell.coordinate,
                                    "refSheet": sheet_ref,
                                    "refCell": m.group(3),
                                }
                            )
                        sample_formulas += 1

        sheets_out.append(
            {
                "name": name,
                "hidden": hidden,
                "sheetState": state,
                "maxRow": max_row,
                "maxColumn": max_col,
                "headerRow1": headers,
                "formulaCellCount": formula_count,
            }
        )

    named_ranges = []
    for defn in wb.defined_names.values():
        named_ranges.append(
            {
                "name": defn.name,
                "attrText": str(defn.attr_text)[:500],
            }
        )

    wb.close()
    return {
        "schemaVersion": 1,
        "appId": app_id,
        "analyzedAtUtc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "sourceXlsx": f"docs/migrations/live/{app_id}.xlsx",
        "gitShaAtAnalysis": git_sha(),
        "sheetCount": len(sheets_out),
        "sheets": sheets_out,
        "namedRangeCount": len(named_ranges),
        "namedRanges": named_ranges,
        "formulaCrossReferencesSample": formula_refs[:500],
        "notes": [
            "Structural inventory only; no cell values beyond row-1 headers.",
            "formulaCrossReferencesSample is capped; formulas preserved in XLSX export.",
        ],
    }


def write_structural_analysis(app_id: str, xlsx_path: Path) -> Path:
    ANALYSIS_DIR.mkdir(parents=True, exist_ok=True)
    struct = analyze_workbook(app_id, xlsx_path)
    out = ANALYSIS_DIR / f"{app_id}.structure.json"
    out.write_text(json.dumps(struct, indent=2) + "\n", encoding="utf-8")
    return out


def print_help() -> None:
    print(
        """DM workbook snapshot tool

Usage:
  .\\docs\\migrations\\tools\\snapshot-dm-workbooks.ps1 [-List] [-WhatIf] [-Force] [-App APP_ID] [-Help]

Examples:
  .\\docs\\migrations\\tools\\snapshot-dm-workbooks.ps1 -List
  .\\docs\\migrations\\tools\\snapshot-dm-workbooks.ps1 -WhatIf
  .\\docs\\migrations\\tools\\snapshot-dm-workbooks.ps1 -App SLG_DM
  .\\docs\\migrations\\tools\\snapshot-dm-workbooks.ps1 -Force

Drive export requires read access. If export fails with 403:
  clasp login --extra-scopes https://www.googleapis.com/auth/drive.readonly
  # or
  python docs/migrations/tools/snapshot_dm_workbooks.py --authorize

Outputs:
  docs/migrations/live/<APP_ID>.xlsx        (gitignored)
  docs/migrations/live/<APP_ID>.meta.json   (gitignored)
  docs/analysis/dm-family/<APP_ID>.structure.json  (tracked, safe structural data)
"""
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(add_help=False)
    parser.add_argument("--list", action="store_true")
    parser.add_argument("--what-if", action="store_true")
    parser.add_argument("--force", action="store_true")
    parser.add_argument("--app", action="append", default=[])
    parser.add_argument("--help", action="store_true")
    parser.add_argument("--authorize", action="store_true")
    parser.add_argument("--analyze-only", action="append", default=[])
    args = parser.parse_args(argv)

    if args.help:
        print_help()
        return 0

    if args.authorize:
        return run_authorize()

    apps = discover_dm_apps()
    if args.list:
        for app in apps:
            clasp_ok = app.clasp_path.is_file()
            print(f"{app.app_id}\t{app.project_path}\tclasp={'yes' if clasp_ok else 'no'}")
        return 0

    selected = apps
    if args.app:
        wanted = {a.upper() for a in args.app}
        selected = [a for a in apps if a.app_id in wanted]
        missing = wanted - {a.app_id for a in selected}
        if missing:
            raise SystemExit(f"Unknown or non-DM app id(s): {', '.join(sorted(missing))}")

    if args.analyze_only:
        for app_id in args.analyze_only:
            xlsx = LIVE_DIR / f"{app_id}.xlsx"
            if not xlsx.is_file():
                print(f"Skip analyze {app_id}: missing {xlsx}", file=sys.stderr)
                continue
            out = write_structural_analysis(app_id, xlsx)
            print(f"Wrote {out.relative_to(REPO_ROOT)}")
        return 0

    clasprc = get_clasprc_credentials()
    script_token = clasprc["access_token"]
    drive_creds = get_drive_credentials()
    drive_token = drive_creds["access_token"]
    freshness = parse_freshness_crosscheck()

    results = []
    for app in selected:
        print(f"=== {app.app_id} ===")
        try:
            res = snapshot_one(
                app,
                what_if=args.what_if,
                force=args.force,
                script_token=script_token,
                drive_token=drive_token,
                drive_auth_source=drive_creds["source"],
                freshness=freshness,
            )
            results.append(res)
            print(json.dumps({k: v for k, v in res.items() if k not in ("freshnessCrossCheck",)}, indent=2))
            if res.get("freshnessCrossCheck"):
                print(res["freshnessCrossCheck"], file=sys.stderr)

            if res.get("action") == "exported":
                xlsx = LIVE_DIR / f"{app.app_id}.xlsx"
                out = write_structural_analysis(app.app_id, xlsx)
                print(f"Structural analysis: {out.relative_to(REPO_ROOT)}")
        except Exception as exc:  # noqa: BLE001
            print(f"ERROR: {exc}", file=sys.stderr)
            results.append({"appId": app.app_id, "action": "failed", "error": str(exc)})

    failed = [r for r in results if r.get("action") == "failed"]
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
