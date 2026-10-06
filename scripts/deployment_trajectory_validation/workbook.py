"""Read-only workbook loading with schema resolution."""

from __future__ import annotations

from typing import Any, Dict, List, Optional, Tuple

import openpyxl

from .schema import (
    SHEET_CANONICAL,
    SchemaResolutionError,
    build_header_map,
    normalize_header,
    resolve_sheet_columns,
)


def resolve_sheet(wb: openpyxl.Workbook, canonical: str) -> Optional[str]:
    names = set(wb.sheetnames)
    for candidate in SHEET_CANONICAL.get(canonical, [canonical]):
        if candidate in names:
            return candidate
    for n in wb.sheetnames:
        if n.startswith(canonical[:20]):
            return n
    return None


def read_sheet_rows(
    wb: openpyxl.Workbook, canonical: str, require_schema: bool = True
) -> Tuple[Optional[str], List[Dict[str, Any]], Dict[str, str]]:
    sheet_name = resolve_sheet(wb, canonical)
    if not sheet_name:
        return None, [], {}
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return sheet_name, [], {}
    headers = [normalize_header(h) for h in header_row]
    header_map = build_header_map(headers)
    column_map: Dict[str, str] = {}
    if require_schema:
        column_map = resolve_sheet_columns(canonical, header_map)

    out: List[Dict[str, Any]] = []
    for row in rows_iter:
        if row is None:
            continue
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        rec: Dict[str, Any] = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            rec[h] = row[i] if i < len(row) else ""
        out.append(rec)
    return sheet_name, out, column_map


def load_workbook_data(
    wb_path: str,
) -> Tuple[Dict[str, Any], Dict[str, List[Dict[str, Any]]], Dict[str, Dict[str, str]]]:
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    sheet_inventory: Dict[str, Any] = {}
    data: Dict[str, List[Dict[str, Any]]] = {}
    column_maps: Dict[str, Dict[str, str]] = {}

    for canonical in SHEET_CANONICAL:
        sheet_name = resolve_sheet(wb, canonical)
        if sheet_name is None:
            sheet_inventory[canonical] = {
                "resolved_name": None,
                "present": False,
                "data_rows": 0,
                "column_count": 0,
            }
            data[canonical] = []
            continue
        try:
            actual, rows, colmap = read_sheet_rows(wb, canonical, require_schema=True)
        except SchemaResolutionError:
            raise
        sheet_inventory[canonical] = {
            "resolved_name": actual,
            "present": actual is not None,
            "data_rows": len(rows),
            "column_count": len(rows[0].keys()) if rows else 0,
            "resolved_columns": colmap,
        }
        data[canonical] = rows
        column_maps[canonical] = colmap

    wb.close()
    return sheet_inventory, data, column_maps
