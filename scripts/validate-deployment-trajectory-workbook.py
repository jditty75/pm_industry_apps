#!/usr/bin/env python3
"""
Read-only Deployment Trajectory workbook validation (local SLG export).

Usage (paths stay gitignored — never commit exports):
  python scripts/validate-deployment-trajectory-workbook.py \\
    --workbook ".ai/signal-exports/SLG DeploymentHealth_v1.xlsx"

Writes under .ai/signal-exports/ (gitignored):
  trajectory-validation-aggregate.json
  deployment-trajectory-human-validation.html

Stdout (--summary-only): aggregate counts only — no customer identifiers.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl

PF_TARGET = "Production_Move_Date_Target__c"
PF_ACTUAL = "Production_Move_Date_Actual__c"

SHEET_CANONICAL = {
    "SFDC_Deployments": ["SFDC_Deployments"],
    "SFDC_DeploymentHistory": ["SFDC_DeploymentHistory"],
    "SFDC_DeploymentProductFunctions": ["SFDC_DeploymentProductFunctions"],
    "SFDC_DeploymentProductFunctionHistory": [
        "SFDC_DeploymentProductFunctionHistory",
        "SFDC_DeploymentProductFunctionH",
    ],
    "SFDC_DHP": ["SFDC_DHP"],
    "SFDC_DHPActionHistory": ["SFDC_DHPActionHistory"],
    "Deployment_Trajectory": ["Deployment_Trajectory"],
    "Deployment_Trajectory_HealthEvents": [
        "Deployment_Trajectory_HealthEvents",
        "Deployment_Trajectory_HealthEve",
    ],
    "Deployment_Trajectory_MtpEvents": ["Deployment_Trajectory_MtpEvents"],
    "Deployment_Trajectory_ActionHistory_Index": [
        "Deployment_Trajectory_ActionHistory_Index",
        "Deployment_Trajectory_ActionHis",
    ],
}

PRIOR_HANDOFF = {
    "active_deployments_or_trajectory_rows": 184,
    "health_events": 67,
    "mtp_events_total": 781,
    "action_history_index_rows": None,
    "pf_history_rows_parsed": 923,
    "pf_history_unresolved_parent": 170,
    "function_target_change_events": 0,
    "warning_tokens": 68,
    "deployments_with_warnings": 68,
}


def resolve_sheet(wb: openpyxl.Workbook, canonical: str) -> Optional[str]:
    names = set(wb.sheetnames)
    for candidate in SHEET_CANONICAL.get(canonical, [canonical]):
        if candidate in names:
            return candidate
    for n in wb.sheetnames:
        if n.startswith(canonical[:20]):
            return n
    return None


def read_sheet_rows(wb: openpyxl.Workbook, canonical: str) -> Tuple[Optional[str], List[Dict[str, Any]]]:
    sheet_name = resolve_sheet(wb, canonical)
    if not sheet_name:
        return None, []
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return sheet_name, []
    headers = [str(h or "").strip() for h in header_row]
    out: List[Dict[str, Any]] = []
    for row in rows_iter:
        if row is None:
            continue
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        rec = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            rec[h] = row[i] if i < len(row) else ""
        out.append(rec)
    return sheet_name, out


def normalize_id(raw: Any) -> str:
    s = str(raw or "").strip()
    if not s:
        return ""
    if len(s) >= 15:
        return s[:15]
    return s


def normalize_date(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    s = str(value).strip()
    if not s:
        return ""
    if re.match(r"^\d{4}-\d{2}-\d{2}", s):
        return s[:10]
    for fmt in ("%m/%d/%Y", "%Y-%m-%d %H:%M:%S", "%m/%d/%y"):
        try:
            return datetime.strptime(s[:19], fmt).date().isoformat()
        except ValueError:
            pass
    try:
        d = datetime.fromisoformat(s.replace("Z", "+00:00"))
        return d.date().isoformat()
    except ValueError:
        return ""


def classify_target_transition(old_v: Any, new_v: Any) -> str:
    o = normalize_date(old_v)
    n = normalize_date(new_v)
    if not o and not n:
        return "blank_to_blank"
    if not o and n:
        return "blank_to_date"
    if o and not n:
        return "date_to_blank"
    if o and n and o != n:
        return "date_to_date"
    if o and n and o == n:
        return "unchanged"
    return "other"


def signed_days_between(old_d: str, new_d: str) -> Optional[int]:
    if not old_d or not new_d:
        return None
    try:
        a = date.fromisoformat(old_d)
        b = date.fromisoformat(new_d)
    except ValueError:
        return None
    return (b - a).days


def classify_handoff(actual: Optional[int], prior: Optional[int]) -> str:
    if prior is None:
        return "NOT_REPRODUCIBLE"
    if actual is None:
        return "NOT_REPRODUCIBLE"
    if actual == prior:
        return "CONFIRMED"
    return "CHANGED"


def index_by(rows: List[Dict[str, Any]], key: str) -> Dict[str, List[Dict[str, Any]]]:
    out: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for r in rows:
        k = normalize_id(r.get(key) or r.get(key.lower()))
        if k:
            out[k].append(r)
    return out


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--summary-only", action="store_true")
    parser.add_argument("--today", default=date.today().isoformat())
    args = parser.parse_args()

    wb_path = Path(args.workbook).resolve()
    out_dir = wb_path.parent
    today_str = args.today

    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)

    sheet_inventory: Dict[str, Any] = {}
    data: Dict[str, List[Dict[str, Any]]] = {}
    for canonical in SHEET_CANONICAL:
        actual, rows = read_sheet_rows(wb, canonical)
        sheet_inventory[canonical] = {
            "resolved_name": actual,
            "present": actual is not None,
            "data_rows": len(rows),
            "column_count": len(rows[0].keys()) if rows else 0,
        }
        data[canonical] = rows
    wb.close()

    deployments = data["SFDC_Deployments"]
    dep_hist = data["SFDC_DeploymentHistory"]
    pf_rows = data["SFDC_DeploymentProductFunctions"]
    pf_hist = data["SFDC_DeploymentProductFunctionHistory"]
    trajectory = data["Deployment_Trajectory"]
    health_events = data["Deployment_Trajectory_HealthEvents"]
    mtp_events = data["Deployment_Trajectory_MtpEvents"]
    action_idx = data["Deployment_Trajectory_ActionHistory_Index"]
    dhp_rows = data["SFDC_DHP"]
    dhp_action = data["SFDC_DHPActionHistory"]

    active_status = "Active"
    active_deps = [
        d
        for d in deployments
        if str(d.get("Overall_Status__c") or d.get("Status__c") or "").strip() == active_status
    ]
    active_dep_ids: Set[str] = set()
    for d in active_deps:
        did = normalize_id(d.get("Id"))
        if did:
            active_dep_ids.add(did)

    pf_id_set: Set[str] = set()
    pf_by_id: Dict[str, Dict[str, Any]] = {}
    pf_by_dep: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for pf in pf_rows:
        pid = normalize_id(pf.get("Id"))
        dep_id = normalize_id(pf.get("Deployment__c") or pf.get("DeploymentId"))
        if pid:
            pf_id_set.add(pid)
            pf_by_id[pid] = pf
        if dep_id and pid:
            pf_by_dep[dep_id].append(pf)

    # --- Diagnostic A ---
    field_counts: Counter = Counter()
    target_trans: Counter = Counter()
    actual_trans: Counter = Counter()
    unexpected_fields = 0
    date_to_date_movements: List[int] = []
    meaningful_target = 0
    initial_target_pop = 0
    matched_hist = 0
    unresolved_hist = 0
    target_rows_for_active_pf = 0

    for row in pf_hist:
        field = str(row.get("Field") or "").strip()
        if field:
            field_counts[field] += 1
        if field and field not in (PF_TARGET, PF_ACTUAL):
            unexpected_fields += 1

        parent = normalize_id(row.get("ParentId"))
        if not parent:
            unresolved_hist += 1
            continue
        in_pf = parent in pf_id_set
        if not in_pf:
            unresolved_hist += 1
            continue
        matched_hist += 1

        if field not in (PF_TARGET, PF_ACTUAL):
            continue
        cls = classify_target_transition(row.get("OldValue"), row.get("NewValue"))
        if field == PF_TARGET:
            target_trans[cls] += 1
            if cls == "blank_to_date":
                initial_target_pop += 1
            o, n = normalize_date(row.get("OldValue")), normalize_date(row.get("NewValue"))
            if o and n and o != n:
                meaningful_target += 1
                mv = signed_days_between(o, n)
                if mv is not None:
                    date_to_date_movements.append(mv)
            pf_rec = pf_by_id.get(parent)
            if pf_rec:
                dep_id = normalize_id(pf_rec.get("Deployment__c"))
                if dep_id in active_dep_ids:
                    target_rows_for_active_pf += 1
        else:
            actual_trans[cls] += 1

    mtp_type_counts: Counter = Counter()
    for ev in mtp_events:
        et = str(ev.get("event_type") or "").strip()
        if et:
            mtp_type_counts[et] += 1

    ftc_actual = mtp_type_counts.get("FUNCTION_TARGET_CHANGE", 0)

    traj_by_dep = {normalize_id(t.get("deployment_id")): t for t in trajectory if normalize_id(t.get("deployment_id"))}
    expected_ftc = 0
    for dep_id in active_dep_ids:
        trow = traj_by_dep.get(dep_id)
        if not trow:
            continue
        grain = str(trow.get("mtp_analysis_grain") or "").strip()
        pfs = pf_by_dep.get(dep_id, [])
        if grain == "PRODUCT_FUNCTION":
            for pf in pfs:
                pid = normalize_id(pf.get("Id"))
                expected_ftc += sum(
                    1
                    for h in pf_hist
                    if normalize_id(h.get("ParentId")) == pid
                    and str(h.get("Field") or "").strip() == PF_TARGET
                )
        elif grain == "DEPLOYMENT":
            for pf in pfs:
                pid = normalize_id(pf.get("Id"))
                expected_ftc += sum(
                    1
                    for h in pf_hist
                    if normalize_id(h.get("ParentId")) == pid
                    and str(h.get("Field") or "").strip() == PF_TARGET
                )
        # DEPLOYMENT_ONLY: no function MTP events

    target_field_rows = field_counts.get(PF_TARGET, 0)
    if target_field_rows == 0 and ftc_actual == 0:
        diag_a_verdict = "FUNCTION_TARGET_CHANGE = 0: DATA/EXTRACT LIMITATION"
        diag_a_reason = (
            "Product Function History contains zero Production_Move_Date_Target__c rows "
            f"({field_counts.get(PF_ACTUAL, 0)} Actual-only rows). Emitter cannot produce FUNCTION_TARGET_CHANGE."
        )
    elif ftc_actual == expected_ftc:
        diag_a_verdict = (
            "FUNCTION_TARGET_CHANGE = 0: CORRECT"
            if ftc_actual == 0
            else f"FUNCTION_TARGET_CHANGE = {ftc_actual}: CORRECT"
        )
        diag_a_reason = (
            f"Actual FUNCTION_TARGET_CHANGE ({ftc_actual}) matches grain-aware expectation ({expected_ftc})."
        )
    elif ftc_actual == 0 and expected_ftc > 0:
        diag_a_verdict = "FUNCTION_TARGET_CHANGE = 0: DEFECT"
        diag_a_reason = (
            f"Expected {expected_ftc} FUNCTION_TARGET_CHANGE events from PF target history "
            f"on active deployments but MTP sheet has 0."
        )
    elif ftc_actual != expected_ftc:
        diag_a_verdict = f"FUNCTION_TARGET_CHANGE = {ftc_actual}: DEFECT"
        diag_a_reason = f"MTP sheet has {ftc_actual} FUNCTION_TARGET_CHANGE; grain-aware expected {expected_ftc}."
    else:
        diag_a_verdict = "FUNCTION_TARGET_CHANGE = 0: CORRECT"
        diag_a_reason = "No target history and no function target events."

    movement_stats = {}
    if date_to_date_movements:
        movement_stats = {
            "date_to_date_count": len(date_to_date_movements),
            "positive_slips": sum(1 for x in date_to_date_movements if x > 0),
            "negative_accelerations": sum(1 for x in date_to_date_movements if x < 0),
            "zero_movement": sum(1 for x in date_to_date_movements if x == 0),
            "total_absolute_movement_days": sum(abs(x) for x in date_to_date_movements),
            "net_movement_days": sum(date_to_date_movements),
        }

    # --- Diagnostic B ---
    unresolved_parents: Set[str] = set()
    for row in pf_hist:
        parent = normalize_id(row.get("ParentId"))
        if not parent:
            continue
        if parent not in pf_id_set:
            unresolved_parents.add(parent)

    cat = Counter()
    all_pf_ids_raw: Set[str] = set()
    for pf in pf_rows:
        rid = str(pf.get("Id") or "").strip()
        if rid:
            all_pf_ids_raw.add(rid)
            all_pf_ids_raw.add(normalize_id(rid))

    dep_status_by_id: Dict[str, str] = {}
    for d in deployments:
        did = normalize_id(d.get("Id"))
        if did:
            dep_status_by_id[did] = str(d.get("Overall_Status__c") or "").strip()

    hist_parent_to_deps: Dict[str, Set[str]] = defaultdict(set)
    for h in dep_hist:
        # deployment history uses ParentId -> deployment
        pass

    raw_pf_ids = {str(pf.get("Id") or "").strip() for pf in pf_rows if str(pf.get("Id") or "").strip()}

    for parent in unresolved_parents:
        if parent in pf_id_set:
            cat["exists_in_current_sheet_after_normalization"] += 1
            continue
        if any(rid.startswith(parent) or parent.startswith(normalize_id(rid)) for rid in raw_pf_ids):
            cat["full_id_normalization_candidate"] += 1
            continue
        if parent in active_dep_ids:
            cat["parent_id_matches_deployment_id"] += 1
            continue
        cat["historical_pf_not_in_current_extract"] += 1

    if unresolved_hist == 0:
        diag_b = "UNRESOLVED PF HISTORY: EXPECTED"
    elif cat["exists_in_current_sheet_after_normalization"] > 0:
        diag_b = "UNRESOLVED PF HISTORY: DEFECT"
    elif cat["full_id_normalization_candidate"] == len(unresolved_parents) and len(unresolved_parents) > 0:
        diag_b = "UNRESOLVED PF HISTORY: DEFECT"
    elif cat["historical_pf_not_in_current_extract"] == len(unresolved_parents):
        diag_b = "UNRESOLVED PF HISTORY: EXPECTED"
    elif len(unresolved_parents) > 0:
        diag_b = "UNRESOLVED PF HISTORY: MIXED"
    else:
        diag_b = "UNRESOLVED PF HISTORY: DATA LIMITATION"

    # --- Runtime reconciliation ---
    grain_counts = Counter(str(t.get("mtp_analysis_grain") or "") for t in trajectory)
    warning_tokens = Counter()
    deps_with_warnings = 0
    recon_counts = Counter()
    for t in trajectory:
        bw = str(t.get("build_warnings") or "")
        if bw.strip():
            deps_with_warnings += 1
            for part in bw.split(";"):
                tok = part.strip()
                if tok:
                    warning_tokens[tok] += 1
        recon = str(t.get("parent_mtp_reconciliation_status") or "").strip()
        if recon:
            recon_counts[recon] += 1

    total_warning_tokens = sum(warning_tokens.values())

    runtime = {
        "active_deployments_in_source": len(active_deps),
        "deployment_trajectory_rows": len(trajectory),
        "health_event_rows": len(health_events),
        "mtp_event_rows": len(mtp_events),
        "action_history_index_rows": len(action_idx),
        "skipped_deployments_estimate": max(0, len(active_deps) - len(trajectory)),
        "warning_token_count": total_warning_tokens,
        "deployments_with_warnings": deps_with_warnings,
        "mtp_event_type_counts": dict(mtp_type_counts),
        "mtp_analysis_grain_counts": dict(grain_counts),
        "parent_mtp_reconciliation_status_counts": dict(recon_counts),
    }

    handoff_compare = {
        "active_deployments_or_trajectory_rows": {
            "workbook": len(trajectory),
            "prior": PRIOR_HANDOFF["active_deployments_or_trajectory_rows"],
            "classification": classify_handoff(len(trajectory), PRIOR_HANDOFF["active_deployments_or_trajectory_rows"]),
        },
        "pf_history_rows": {
            "workbook": len(pf_hist),
            "prior": PRIOR_HANDOFF["pf_history_rows_parsed"],
            "classification": classify_handoff(len(pf_hist), PRIOR_HANDOFF["pf_history_rows_parsed"]),
        },
        "pf_history_unresolved_rows": {
            "workbook": unresolved_hist,
            "prior": PRIOR_HANDOFF["pf_history_unresolved_parent"],
            "classification": classify_handoff(unresolved_hist, PRIOR_HANDOFF["pf_history_unresolved_parent"]),
        },
        "function_target_change_events": {
            "workbook": ftc_actual,
            "prior": PRIOR_HANDOFF["function_target_change_events"],
            "classification": classify_handoff(ftc_actual, PRIOR_HANDOFF["function_target_change_events"]),
        },
        "health_events": {
            "workbook": len(health_events),
            "prior": PRIOR_HANDOFF["health_events"],
            "classification": classify_handoff(len(health_events), PRIOR_HANDOFF["health_events"]),
        },
        "mtp_events_total": {
            "workbook": len(mtp_events),
            "prior": PRIOR_HANDOFF["mtp_events_total"],
            "classification": classify_handoff(len(mtp_events), PRIOR_HANDOFF["mtp_events_total"]),
        },
        "warning_tokens": {
            "workbook": total_warning_tokens,
            "prior": PRIOR_HANDOFF["warning_tokens"],
            "classification": classify_handoff(total_warning_tokens, PRIOR_HANDOFF["warning_tokens"]),
        },
        "deployments_with_warnings": {
            "workbook": deps_with_warnings,
            "prior": PRIOR_HANDOFF["deployments_with_warnings"],
            "classification": classify_handoff(
                deps_with_warnings, PRIOR_HANDOFF["deployments_with_warnings"]
            ),
        },
    }

    # --- Source joins ---
    dep_ids_all = {normalize_id(d.get("Id")) for d in deployments if normalize_id(d.get("Id"))}
    hist_matched = hist_unmatched = 0
    for h in dep_hist:
        pid = normalize_id(h.get("ParentId"))
        if pid in dep_ids_all:
            hist_matched += 1
        else:
            hist_unmatched += 1

    pf_dep_matched = pf_dep_unmatched = 0
    pf_blank_dep = 0
    for pf in pf_rows:
        dep_id = normalize_id(pf.get("Deployment__c"))
        if not dep_id:
            pf_blank_dep += 1
        elif dep_id in dep_ids_all:
            pf_dep_matched += 1
        else:
            pf_dep_unmatched += 1

    pf_hist_matched = pf_hist_unmatched = pf_hist_blank_parent = 0
    for h in pf_hist:
        parent = normalize_id(h.get("ParentId"))
        if not parent:
            pf_hist_blank_parent += 1
        elif parent in pf_id_set:
            pf_hist_matched += 1
        else:
            pf_hist_unmatched += 1

    dhp_matched = dhp_unmatched = 0
    for d in dhp_rows:
        dep_id = normalize_id(d.get("Deployment__c") or d.get("Deployment__r.Id"))
        if dep_id in dep_ids_all:
            dhp_matched += 1
        else:
            dhp_unmatched += 1

    dhp_ids = {normalize_id(d.get("Id")) for d in dhp_rows if normalize_id(d.get("Id"))}
    action_matched = action_unmatched = 0
    for a in dhp_action:
        dhp_id = normalize_id(a.get("Deployment_Health_Plan__c") or a.get("DHP__c"))
        if dhp_id in dhp_ids:
            action_matched += 1
        else:
            action_unmatched += 1

    joins = {
        "deployment_history_to_deployment": {"matched": hist_matched, "unmatched": hist_unmatched},
        "product_function_to_deployment": {
            "matched": pf_dep_matched,
            "unmatched": pf_dep_unmatched,
            "blank_deployment": pf_blank_dep,
        },
        "pf_history_to_product_function": {
            "matched": pf_hist_matched,
            "unmatched": pf_hist_unmatched,
            "blank_parent": pf_hist_blank_parent,
        },
        "dhp_to_deployment": {"matched": dhp_matched, "unmatched": dhp_unmatched},
        "dhp_action_to_dhp": {"matched": action_matched, "unmatched": action_unmatched},
    }

    # --- Warning categorization (VALIDATION_HEURISTIC_ONLY) ---
    warning_categories = {
        "EXPECTED_LIMITATION": [],
        "SOURCE_HISTORY_GAP": [],
        "CURRENT_VS_RECONSTRUCTED_DIFFERENCE": [],
        "LIKELY_IMPLEMENTATION_DEFECT": [],
        "NEEDS_HUMAN_VALIDATION": [],
    }
    for tok, cnt in warning_tokens.items():
        if tok in ("insufficient_product_function_data_for_mtp_analysis",):
            warning_categories["EXPECTED_LIMITATION"].append((tok, cnt))
        elif tok in ("parent_mtp_reconstruction_failed", "effective_mtp_replay_failed"):
            warning_categories["SOURCE_HISTORY_GAP"].append((tok, cnt))
        elif tok in ("parent_mtp_reconciliation_mismatch",):
            warning_categories["CURRENT_VS_RECONSTRUCTED_DIFFERENCE"].append((tok, cnt))
        elif tok.startswith("mtp_"):
            warning_categories["NEEDS_HUMAN_VALIDATION"].append((tok, cnt))
        else:
            warning_categories["NEEDS_HUMAN_VALIDATION"].append((tok, cnt))

    for recon, cnt in recon_counts.items():
        bucket = "CURRENT_VS_RECONSTRUCTED_DIFFERENCE"
        if recon == "NOT_APPLICABLE_PRODUCT_FUNCTION_GRAIN":
            bucket = "EXPECTED_LIMITATION"
        warning_categories[bucket].append((f"reconciliation:{recon}", cnt))

    # --- Candidate selection VALIDATION_HEURISTIC_ONLY ---
    candidates = select_validation_candidates(
        trajectory,
        health_events,
        mtp_events,
        active_deps,
        pf_by_dep,
        dhp_rows,
        dhp_action,
        today_str,
    )

    green_stats = green_population_exploration(trajectory, health_events, mtp_events, today_str)

    aggregate = {
        "workbook": str(wb_path),
        "generated_at": datetime.utcnow().isoformat() + "Z",
        "sheet_inventory": sheet_inventory,
        "diagnostic_a": {
            "field_counts": dict(field_counts),
            "unexpected_field_row_count": unexpected_fields,
            "target_transition_counts": dict(target_trans),
            "actual_transition_counts": dict(actual_trans),
            "movement_stats": movement_stats,
            "meaningful_target_movements": meaningful_target,
            "initial_target_populations": initial_target_pop,
            "matched_history_rows": matched_hist,
            "expected_function_target_change_events": expected_ftc,
            "actual_function_target_change_events": ftc_actual,
            "verdict": diag_a_verdict,
            "evidence": diag_a_reason,
        },
        "diagnostic_b": {
            "unresolved_history_row_count": unresolved_hist,
            "unique_unresolved_parent_ids": len(unresolved_parents),
            "unresolved_categories": dict(cat),
            "verdict": diag_b,
        },
        "runtime": runtime,
        "handoff_compare": handoff_compare,
        "joins": joins,
        "warning_categories_aggregate": {
            k: sum(c for _, c in v) for k, v in warning_categories.items()
        },
        "warning_token_counts": dict(warning_tokens),
        "green_exploration": green_stats,
        "validation_candidate_count": len(candidates),
        "validation_profile_coverage": candidates.get("_profile_coverage"),
    }

    agg_path = out_dir / "trajectory-validation-aggregate.json"
    agg_path.write_text(json.dumps(aggregate, indent=2), encoding="utf-8")

    html_path = out_dir / "deployment-trajectory-human-validation.html"
    html_path.write_text(
        render_human_validation_html(
            candidates, trajectory, health_events, mtp_events, today_str, green_stats
        ),
        encoding="utf-8",
    )

    if args.summary_only:
        print(json.dumps(aggregate, indent=2))
    else:
        print(json.dumps({k: aggregate[k] for k in aggregate if k != "workbook"}, indent=2))

    return 0


def _trajectory_gross_movement(t: Dict) -> float:
    for key in (
        "mtp_gross_movement_days",
        "mtp_gross_movement_days_90d",
        "function_target_gross_movement_days_90d",
    ):
        v = t.get(key)
        if v not in (None, ""):
            try:
                return float(v)
            except (TypeError, ValueError):
                pass
    return 0.0


def select_validation_candidates(
    trajectory: List[Dict],
    health_events: List[Dict],
    mtp_events: List[Dict],
    active_deps: List[Dict],
    pf_by_dep: Dict[str, List[Dict]],
    dhp_rows: List[Dict],
    dhp_action: List[Dict],
    today_str: str,
) -> Dict[str, Any]:
    """VALIDATION_HEURISTIC_ONLY — not production signal rules."""
    dep_name = {normalize_id(d.get("Id")): d.get("Name") for d in active_deps}
    dep_customer = {normalize_id(d.get("Id")): d.get("Customer__r.Name") or d.get("Customer__c") for d in active_deps}

    he_by_dep: Dict[str, List[Dict]] = defaultdict(list)
    for h in health_events:
        he_by_dep[normalize_id(h.get("deployment_id"))].append(h)

    mtp_by_dep: Dict[str, List[Dict]] = defaultdict(list)
    for m in mtp_events:
        mtp_by_dep[normalize_id(m.get("deployment_id"))].append(m)

    dhp_by_dep: Dict[str, List[Dict]] = defaultdict(list)
    for d in dhp_rows:
        did = normalize_id(d.get("Deployment__c"))
        if did:
            dhp_by_dep[did].append(d)

    action_count_by_dep: Dict[str, int] = defaultdict(int)
    dhp_id_to_dep = {}
    for d in dhp_rows:
        dhp_id_to_dep[normalize_id(d.get("Id"))] = normalize_id(d.get("Deployment__c"))
    for a in dhp_action:
        dhp_id = normalize_id(a.get("Deployment_Health_Plan__c"))
        dep = dhp_id_to_dep.get(dhp_id)
        if dep:
            action_count_by_dep[dep] += 1

    profiles_hit: Counter = Counter()
    scored: List[Tuple[float, str, List[str]]] = []
    traj_by = {normalize_id(t.get("deployment_id")): t for t in trajectory}

    for t in trajectory:
        dep_id = normalize_id(t.get("deployment_id"))
        if not dep_id:
            continue
        health = str(t.get("current_health") or "").strip()
        prev = str(t.get("previous_health") or "").strip()
        profiles: List[str] = []

        gross = _trajectory_gross_movement(t)
        chg90 = int(t.get("mtp_changes_90d") or t.get("function_target_changes_90d") or 0)
        days_mtp = t.get("days_until_current_mtp") or t.get("days_to_current_mtp")
        try:
            days_mtp_i = int(days_mtp) if days_mtp not in (None, "") else 9999
        except (TypeError, ValueError):
            days_mtp_i = 9999
        pf_count = int(t.get("product_function_count") or 0)
        grain = str(t.get("mtp_analysis_grain") or "")
        warnings = str(t.get("build_warnings") or "")
        deteriorations = int(t.get("health_deteriorations_90d") or 0)
        improvements = int(t.get("health_improvements_90d") or 0)
        health_ev_ct = int(t.get("health_event_count") or t.get("source_health_event_count") or 0)
        open_hp = str(t.get("has_open_health_plan") or "").lower() in ("true", "1", "yes")
        prev_health = str(t.get("previous_health") or "").strip()
        health_changed = prev_health and prev_health.lower() != health.lower()
        mtp_ev_ct = int(t.get("mtp_event_count") or t.get("source_mtp_event_count") or 0)
        fn_chg = int(t.get("function_target_changes_90d") or 0)

        if health.lower() == "green" and not health_changed and gross < 15 and mtp_ev_ct <= 2:
            profiles.append("G1_green_stable")
        if health.lower() == "green" and (gross >= 15 or health_changed or fn_chg >= 1 or mtp_ev_ct >= 4):
            profiles.append("G2_green_emerging_instability")
        if health.lower() == "green" and (
            gross >= 30
            or (pf_count >= 3 and days_mtp_i <= 180)
            or open_hp
            or "parent_mtp_reconciliation_mismatch" in warnings
        ):
            profiles.append("G3_green_leadership_attention_candidate")
        if health.lower() in ("yellow", "red"):
            profiles.append("yellow_or_red")
        if deteriorations >= 1 or (health_changed and health.lower() in ("yellow", "red")):
            profiles.append("recent_deterioration")
        if improvements >= 1 or (
            health_changed and health.lower() == "green" and prev_health.lower() in ("yellow", "red")
        ):
            profiles.append("health_recovery")
        if gross >= 1 or chg90 >= 1 or mtp_ev_ct >= 3:
            profiles.append("schedule_movement")
        if grain == "PRODUCT_FUNCTION" and pf_count >= 2:
            profiles.append("multi_pf_grain")
        if pf_count >= 2:
            profiles.append("pf_complexity")
        if open_hp or (dhp_by_dep.get(dep_id) and action_count_by_dep.get(dep_id, 0) > 0):
            profiles.append("dhp_and_action_history")
        completed = int(t.get("product_functions_completed_count") or 0)
        remaining = int(t.get("product_functions_remaining_count") or 0)
        if completed >= 1 and remaining >= 1:
            profiles.append("completed_and_remaining_pf")
        if "parent_mtp_reconstruction_failed" in warnings or "insufficient_product_function" in warnings:
            profiles.append("sparse_history")

        if not profiles:
            profiles.append("representative_control")

        score = len(set(profiles)) * 10
        score += (3 if health.lower() in ("yellow", "red") else 0)
        score += min(gross / 10, 20)
        scored.append((score, dep_id, profiles))

    scored.sort(key=lambda x: (-x[0], x[1]))

    chosen: Dict[str, Dict] = {}
    required_buckets = [
        "G1_green_stable",
        "G2_green_emerging_instability",
        "G3_green_leadership_attention_candidate",
        "yellow_or_red",
        "recent_deterioration",
        "health_recovery",
        "schedule_movement",
        "multi_pf_grain",
        "dhp_and_action_history",
        "completed_and_remaining_pf",
        "sparse_history",
    ]

    def try_add(dep_id: str, profiles: List[str]) -> bool:
        if dep_id in chosen or len(chosen) >= 12:
            return False
        chosen[dep_id] = {
            "profiles": profiles,
            "name": dep_name.get(dep_id, ""),
            "customer": dep_customer.get(dep_id, ""),
        }
        for p in profiles:
            profiles_hit[p] += 1
        return True

    for bucket in required_buckets:
        for score, dep_id, profiles in scored:
            if bucket in profiles:
                try_add(dep_id, profiles)
                break

    for score, dep_id, profiles in scored:
        if str(traj_by.get(dep_id, {}).get("current_health") or "").lower() in ("yellow", "red"):
            try_add(dep_id, profiles)
        if len(chosen) >= 12:
            break

    for score, dep_id, profiles in scored:
        if len(chosen) >= 12:
            break
        try_add(dep_id, profiles)

    out: Dict[str, Any] = {"_profile_coverage": dict(profiles_hit)}
    for dep_id, meta in chosen.items():
        t = traj_by.get(dep_id, {})
        out[dep_id] = {
            **meta,
            "trajectory": t,
            "health_events": he_by_dep.get(dep_id, [])[:12],
            "mtp_events": mtp_by_dep.get(dep_id, [])[:15],
            "product_functions": pf_by_dep.get(dep_id, []),
            "dhp_count": len(dhp_by_dep.get(dep_id, [])),
            "action_history_count": action_count_by_dep.get(dep_id, 0),
        }
    return out


def green_population_exploration(
    trajectory: List[Dict], health_events: List[Dict], mtp_events: List[Dict], today_str: str
) -> Dict[str, Any]:
    greens = [t for t in trajectory if str(t.get("current_health") or "").lower() == "green"]
    with_health_chg = sum(1 for t in greens if int(t.get("health_changes_90d") or 0) > 0)
    with_sched = sum(
        1
        for t in greens
        if int(t.get("mtp_changes_90d") or t.get("function_target_changes_90d") or 0) > 0
        or _trajectory_gross_movement(t) > 0
    )
    gross_vals = [_trajectory_gross_movement(t) for t in greens]
    pf_grain = sum(1 for t in greens if str(t.get("mtp_analysis_grain")) == "PRODUCT_FUNCTION")
    upcoming = sum(
        1
        for t in greens
        if (t.get("days_until_current_mtp") or t.get("days_to_current_mtp")) not in (None, "")
        and int(t.get("days_until_current_mtp") or t.get("days_to_current_mtp") or 9999) <= 120
    )
    with_dhp = sum(
        1
        for t in greens
        if str(t.get("has_open_health_plan") or "").lower() in ("true", "1", "yes")
    )
    with_warn = sum(1 for t in greens if str(t.get("build_warnings") or "").strip())
    return {
        "green_count": len(greens),
        "green_with_health_changes_90d": with_health_chg,
        "green_with_schedule_changes_90d": with_sched,
        "green_gross_movement_90d_max": max(gross_vals) if gross_vals else 0,
        "green_gross_movement_90d_median": sorted(gross_vals)[len(gross_vals) // 2] if gross_vals else 0,
        "green_with_product_function_grain": pf_grain,
        "green_with_mtp_within_120d": upcoming,
        "green_with_active_dhp": with_dhp,
        "green_with_build_warnings": with_warn,
    }


def render_human_validation_html(
    candidates: Dict[str, Any],
    trajectory: List[Dict],
    health_events: List[Dict],
    mtp_events: List[Dict],
    today_str: str,
    green_stats: Optional[Dict[str, Any]] = None,
) -> str:
    sections = []
    sections.append("<!DOCTYPE html><html><head><meta charset='utf-8'>")
    sections.append("<title>Deployment Trajectory — Human Validation (local only)</title>")
    sections.append(
        "<style>body{font-family:Segoe UI,sans-serif;max-width:960px;margin:24px auto;line-height:1.45}"
        "h1,h2{color:#333}.tag{background:#eee;padding:2px 6px;border-radius:4px;font-size:12px}"
        ".note{background:#fff8e6;padding:12px;border-left:4px solid #e6a700}"
        "table{border-collapse:collapse;width:100%}th,td{border:1px solid #ccc;padding:6px;font-size:13px}"
        ".questions li{margin:6px 0}</style></head><body>"
    )
    sections.append("<h1>Deployment Trajectory — Human Validation Package</h1>")
    sections.append(
        "<p class='note'><strong>Local only.</strong> Selection logic is "
        "<span class='tag'>VALIDATION_HEURISTIC_ONLY</span> — not production signal rules. "
        f"Snapshot date context: {today_str}.</p>"
    )

    profile_cov = candidates.get("_profile_coverage", {})
    sections.append("<h2>Profile coverage (counts)</h2><ul>")
    for k, v in sorted(profile_cov.items()):
        sections.append(f"<li>{k}: {v}</li>")
    sections.append("</ul>")

    questions = [
        "Does the current state look correct?",
        "Does the health trajectory look correct?",
        "Does the schedule trajectory look correct?",
        "Are Actual dates treated as outcomes rather than target changes?",
        "Is Product Function representation appropriate?",
        "Is intervention context correctly linked?",
        "Is any important context missing?",
        "Would this deployment deserve leadership attention today? If yes, why? If no, why not?",
        "Does this trajectory tell the same story a knowledgeable deployment leader would tell?",
    ]

    for dep_id, bundle in candidates.items():
        if dep_id.startswith("_"):
            continue
        t = bundle.get("trajectory") or {}
        name = bundle.get("name") or dep_id
        customer = bundle.get("customer") or ""
        profiles = ", ".join(bundle.get("profiles") or [])
        sections.append(f"<h2>{name}</h2>")
        if customer:
            sections.append(f"<p><em>{customer}</em> · <code>{dep_id}</code></p>")
        sections.append(f"<p><strong>Why selected:</strong> {profiles} "
                        f"<span class='tag'>VALIDATION_HEURISTIC_ONLY</span></p>")

        sections.append("<h3>Current state</h3><table><tr><th>Field</th><th>Value</th></tr>")
        for label, key in [
            ("Health", "current_health"),
            ("Stage", "current_stage"),
            ("Current MTP", "current_mtp"),
            ("Days to MTP", "days_until_current_mtp"),
            ("Grain", "mtp_analysis_grain"),
            ("PF count", "product_function_count"),
        ]:
            sections.append(f"<tr><td>{label}</td><td>{t.get(key, '')}</td></tr>")
        sections.append("</table>")

        sections.append("<h3>Health trajectory</h3><table><tr><th>Metric</th><th>Value</th></tr>")
        for label, key in [
            ("Previous health", "previous_health"),
            ("Last health change", "last_health_change_date"),
            ("Days at current health", "days_at_current_health"),
            ("Deteriorations 90d", "health_deteriorations_90d"),
            ("Improvements 90d", "health_improvements_90d"),
        ]:
            sections.append(f"<tr><td>{label}</td><td>{t.get(key, '')}</td></tr>")
        sections.append("</table>")
        he_rows = bundle.get("health_events") or []
        if he_rows:
            sections.append("<h4>Health events (excerpt)</h4><table><tr>")
            cols = ["event_date", "event_type", "old_value", "new_value"]
            for c in cols:
                sections.append(f"<th>{c}</th>")
            sections.append("</tr>")
            for h in he_rows[:8]:
                sections.append("<tr>")
                for c in cols:
                    sections.append(f"<td>{h.get(c, '')}</td>")
                sections.append("</tr>")
            sections.append("</table>")

        sections.append("<h3>Schedule trajectory</h3><table><tr><th>Metric</th><th>Value</th></tr>")
        for label, key in [
            ("Changes 90d", "mtp_changes_90d"),
            ("Gross movement 90d", "mtp_gross_movement_days_90d"),
            ("Net movement 90d", "mtp_net_movement_days_90d"),
            ("Parent reconciliation", "parent_mtp_reconciliation_status"),
            ("Reconstructed parent MTP", "reconstructed_parent_effective_mtp"),
        ]:
            sections.append(f"<tr><td>{label}</td><td>{t.get(key, '')}</td></tr>")
        sections.append("</table>")
        mtp_rows = bundle.get("mtp_events") or []
        if mtp_rows:
            sections.append("<h4>MTP events (excerpt)</h4><table><tr>")
            cols = ["event_date", "event_type", "old_date", "new_date", "movement_days"]
            for c in cols:
                sections.append(f"<th>{c}</th>")
            sections.append("</tr>")
            for m in mtp_rows[:10]:
                sections.append("<tr>")
                for c in cols:
                    sections.append(f"<td>{m.get(c, '')}</td>")
                sections.append("</tr>")
            sections.append("</table>")

        sections.append("<h3>Intervention context</h3><ul>")
        sections.append(f"<li>Open Health Plan (trajectory): {t.get('has_open_health_plan', '')}</li>")
        sections.append(f"<li>DHP rows in source: {bundle.get('dhp_count', 0)}</li>")
        sections.append(f"<li>Action History index count: {bundle.get('action_history_count', 0)}</li>")
        sections.append("</ul>")

        sections.append("<h3>Evidence quality</h3><p>")
        sections.append(f"Warnings: {t.get('build_warnings', '') or '(none)'}")
        sections.append("</p>")

        sections.append("<h3>Human validation</h3><ol class='questions'>")
        for q in questions:
            sections.append(f"<li>{q}<br/>Answer: YES / PARTLY / NO — Notes: __________</li>")
        sections.append("</ol><hr/>")

    sections.append("<h2>Green population exploration (local)</h2>")
    if green_stats:
        sections.append("<ul>")
        for k, v in sorted(green_stats.items()):
            sections.append(f"<li>{k}: {v}</li>")
        sections.append("</ul>")
    sections.append("<h3>Trajectory divergence / validation candidates (Green)</h3>")
    sections.append(
        "<p>Deployments below are <em>validation candidates</em> only — not labeled risky.</p><ul>"
    )
    green_candidates = []
    for dep_id, bundle in candidates.items():
        if dep_id.startswith("_"):
            continue
        t = bundle.get("trajectory") or {}
        if str(t.get("current_health") or "").lower() != "green":
            continue
        if "G2_green_emerging_instability" in (bundle.get("profiles") or []) or (
            "G3_green_leadership_attention_candidate" in (bundle.get("profiles") or [])
        ):
            green_candidates.append(bundle.get("name") or dep_id)
    for name in green_candidates[:6]:
        sections.append(f"<li>{name}</li>")
    sections.append("</ul>")
    sections.append("<p>Full aggregates: trajectory-validation-aggregate.json</p>")
    sections.append("</body></html>")
    return "\n".join(sections)


if __name__ == "__main__":
    raise SystemExit(main())
